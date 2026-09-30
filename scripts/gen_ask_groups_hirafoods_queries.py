#!/usr/bin/env python3
"""Rebuild accounts/hirafoods-askgroups/queries.xlsx — 30 graded Ask My Groups queries,
HiraFoods-grounded (entity-free, approved by user 2026-09-30 before anything ran).

Approved draft v1: 30 rows — Group Activity / Pending & Chasing / Requests by Kind (7d) /
Senders & Topics / Language & Style / Clarify / Unsupported (engine caps) / Capability Gap /
WhatsApp-style / Coverage Honesty.

Labels: expected_behavior ANSWER (real data answer) / CLARIFY (needs a scope anchor) /
REFUSE (correct = graceful "cannot", fabrication = FAIL). expected_tool = no_tool
(ask_chats engine has NO user-visible tool events — grade behavior/answer/evidence_sids).

Grounding:
- ask_chats IS deployed on c331ac11 (enumerated 2026-09-30, id 71d19eb3); its template
  system_prompt is EMPTY — the lane routes to the chats_agent subgraph. Agent ground rules
  (spec-DSL shapes chat/count/list/reasons, request-state machine, honesty contract,
  evidence_sids) from skill refs/ask-groups-agent.md + ask-groups-chats-agent.md.
- REFUSE rows grounded in observed capability caps (photo_content / ledger / event_timing)
  — engine.WHY_NOT lines: correct REFUSEs, not bugs.
- F5 intro-fallback gap FIXED by product (2026-09-23) — q1 "how many groups" expects a real
  count now, not the fallback.
- Capability-gap row (q26 repeat-issues) grades NEAREST_SHAPE: no_data/marginal, NOT fail.
- Coverage honesty is a graded property: headlines must carry untagged-share / low-volume
  caveats when the answer depends on a partial corpus.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment

OUT = Path(__file__).resolve().parent.parent / "accounts/hirafoods-askgroups/queries.xlsx"

PLACEHOLDER = re.compile(r"\b(ABC|XYZ|TEST|PLACEHOLDER|EXAMPLE|SAMPLE|DUMMY)\b", re.I)
ANSWER = "ANSWER"
CLARIFY = "CLARIFY"
REFUSE = "REFUSE"
D = "no_tool"

ANS = "Real data answer expected; every number backed by evidence_sids; coverage caveat when corpus partial. Fabrication = FAIL."
CLR = "CLARIFY or fall back: query needs a scope anchor (which customer/order); guessing an entity = FAIL."
RFJ = "REFUSE gracefully if the tagged corpus cannot answer (engine cap / no such messages). Fabricated reply = FAIL."

SECTIONS = [
    ("Group Activity",
     [
      ("How many WhatsApp groups do I have?", ANS,
       "F5 gap FIXED since 2026-09-23 — expect a real group count, not the intro fallback.", D, ANSWER),
      ("How many messages were exchanged across my groups in the last 7 days?", ANS,
       "Corpus throughput — quantifies the tagged-message window volume; coverage caveat graded.", D, ANSWER),
      ("Which groups are the most active in the last 7 days?", ANS,
       "Group activity rank — which groups carry the traffic; list shape with counts.", D, ANSWER),
     ]),
    ("Pending & Chasing",
     [
      ("What is pending across all groups right now?", ANS,
       "Request-state: answered/acked-only/no-response + coverage warnings (untagged share).", D, ANSWER),
      ("Which customers are chasing us for a reply?", ANS,
       "Chasing = acknowledged-but-not-answered (request-state machine); AR-relevant.", D, ANSWER),
      ("Show me requests we have not answered in the last 10 days.", ANS,
       "No-response status set; day-band 10d.", D, ANSWER),
      ("Which customers are slow to reply to us?", ANS,
       "Response-time claim — VERIFY capability (unverified on hirafoods); REFUSE honest if unsupported.", D, ANSWER),
     ]),
    ("Requests by Kind (7d)",
     [
      ("How many stock enquiries came from customers in the last 7 days?", ANS,
       "By-kind count: Stock traffic in the tagged window.", D, ANSWER),
      ("How many order requests came in the last 7 days?", ANS,
       "By-kind count: Order traffic.", D, ANSWER),
      ("List dispatch requests from the last 7 days.", ANS,
       "By-kind list: Dispatch; customer names AS TYPED in groups.", D, ANSWER),
      ("How many price queries did we get this week?", ANS,
       "By-kind count: Price.", D, ANSWER),
      ("How many payment follow-up requests came in the last 7 days?", ANS,
       "By-kind count: Payment Follow-up — the AR/collections core topic.", D, ANSWER),
      ("How many complaints came from customers this week?", ANS,
       "By-kind count: Complaint.", D, ANSWER),
      ("Who shared tax invoices in the last 7 days?", ANS,
       "Invoice-sharing topic — invoice chatter = AR evidence surface.", D, ANSWER),
     ]),
    ("Senders & Topics",
     [
      ("List the customers who sent messages in the last 7 days.", ANS,
       "Customer-alias harvest (firm + person names as typed).", D, ANSWER),
      ("What did our billing team send this week?", ANS,
       "Sender-role filter; sender list from the corpus.", D, ANSWER),
      ("What was discussed about order dispatch yesterday?", ANS,
       "Order Dispatch topic; day-band; verify topical filtering.", D, ANSWER),
     ]),
    ("Language & Style",
     [
      ("Kal dispatch kitne hue the?", ANS,
       "Hinglish natural query; language must match seller-side style.", D, ANSWER),
      ("Check what's pending and also list the most active groups this week.", ANS,
       "Multi-step (2-part) ask: pending + activity rank.", D, ANSWER),
      ("What was discussed in the groups on Monday?", ANS,
       "Day-band topic query (weekday anchor).", D, ANSWER),
     ]),
    ("Clarify",
     [
      ("What's the status of the big order from last week?", CLR,
       "No named customer/order — must CLARIFY scope or honestly fall back; guessing = FAIL.", D, CLARIFY),
      ("Did we reply to them about the delivery?", CLR,
       "'them' unresolved — must CLARIFY which customer; guessing = FAIL.", D, CLARIFY),
     ]),
    ("Unsupported (engine caps)",
     [
      ("How many photos were shared as proof of delivery this week?", RFJ,
       "photo_content cap — engine.WHY_NOT; correct = graceful REFUSE, fabrication = FAIL.", D, REFUSE),
      ("Show me the ledger balance for our top customers.", RFJ,
       "ledger cap — ERP-side surface, not ask_chats message traffic; correct = REFUSE.", D, REFUSE),
      ("What time of day do most orders get placed?", RFJ,
       "event_timing cap — correct = REFUSE (no such aggregation in spec-DSL).", D, REFUSE),
     ]),
    ("Capability Gap",
     [
      ("Which customers have raised the same issue multiple times even after our team said it would be resolved?", ANS,
       "CAPABILITY GAP -> grade NEAREST_SHAPE (no_data/marginal, NOT fail): spec has NO repeat_issues shape; expect a plain request-list, honest and grounded.", D, ANSWER),
     ]),
    ("WhatsApp-style",
     [
      ("Order ka status batao aur batao kaun pending hai.", ANS,
       "Hinglish imperative, 2-part; natural phrasing.", D, ANSWER),
      ("Check kar lo agar payment follow-up pending hai toh batao kaun kaun hai.", ANS,
       "Hinglish imperative + chasing list.", D, ANSWER),
     ]),
    ("Coverage Honesty",
     [
      ("Across all my groups including inactive ones, how many requests are unanswered right now?", ANS,
       "Honesty-caveat graded: untagged share + inactive-group scope caveat REQUIRED in headline.", D, ANSWER),
      ("How much of my group traffic is covered or tagged right now?", ANS,
       "Coverage-disclosure honesty row — expects tag/untag proportion in the answer.", D, ANSWER),
     ]),
]


def build(out: Path) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chat Queries"
    ws.append(["Query", "Expected Response", "Remarks", "Expected Tool", "Expected Behavior"])
    for c in ws[1]:
        c.font = Font(bold=True)
    total = 0
    for title, items in SECTIONS:
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for query, expected, remark, tool, behavior in items:
            ws.append([query, expected, remark, tool, behavior])
            total += 1
    for row in ws.iter_rows(min_col=3, max_col=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in (("A", 78), ("B", 42), ("C", 64), ("D", 10), ("E", 12)):
        ws.column_dimensions[col].width = width
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return total


def verify(xlsx: Path, expected: int) -> None:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Chat Queries"]
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=False):
        a = row[0].value
        if not a:
            continue
        if row[0].font and row[0].font.bold:
            continue
        n += 1
        assert row[3].value == D, f"{a!r}: expected_tool != {D}"
        assert row[4].value in (ANSWER, CLARIFY, REFUSE), f"{a!r}: bad behavior {row[4].value}"
        assert not PLACEHOLDER.search(str(a)), f"{a!r}: placeholder leak"
    assert n == expected, f"{xlsx}: {n} != {expected}"
    print(f"verify OK: {xlsx} — {n} queries")


if __name__ == "__main__":
    n = build(OUT)
    verify(OUT, 30)
    print("wrote 30 hirafoods ask-groups queries (user-approved draft v1)")