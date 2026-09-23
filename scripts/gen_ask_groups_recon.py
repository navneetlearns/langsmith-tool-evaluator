#!/usr/bin/env python3
"""Rebuild accounts/ask-groups/queries.xlsx — RECON PROBE set, NOT a graded eval.

Purpose (2026-09-23, user-approved sequence): run a SMALL ask-groups probe on the Zainab
workspace (d53279c2) BEFORE building collections/AR query sets, so we know WHAT EXISTS in the
WhatsApp groups (corpus volume, topics, request statuses, coverage) and harvest real customer
aliases. Findings feed AR/collections query building + the AR entity harvest.

Grading: probes are NOT labeled/graded. expected_tool="RECON" (the chats engine has NO
user-visible tool events — grade behavior/answers+evidence_sids, never tools), and
expected_behavior="RECON" so downstream naıve graders skip the set.

Probe mix is built from the trace-verified surface (agent profile ask-groups-agent.md):
- [T] = trace-observed working: keyword+day-band lists, reasons follow-ups, evidence contract.
- [S] = agent-claimed, unverified: pending counts, chasing, response times -> VERIFY here.
- known gap repro: "how many groups i have" (trace 1 fallback, 570 in context but unused).
- AR-relevant topics: Payment Follow-up, Tax Invoice sharing, Order Dispatch, Stock Enquiry.
"""
import re
import sys

import openpyxl
from openpyxl.styles import Font
from pathlib import Path

OUT = Path("/home/sumit/AgentWork/eval-dashboard/accounts/ask-groups/queries.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

PLACEHOLDER = re.compile(r"(ABC|XYZ|TEST|PLACEHOLDER|EXAMPLE|SAMPLE|DUMMY)", re.I)

# (category, [(query, remark), ...])  — one section: probes run in this order.
SECTIONS = {
    "Recon Probes": [
        ("How many groups do I have?",
         "[T-gap] trace-1 repro: falls back to intro though 570 groups sit in parse memory"),
        ("How many messages were exchanged across my groups in the last 7 days?",
         "corpus throughput — quantifies the tagged-message window volume"),
        ("What is pending across all groups right now?",
         "[S] pending/requests + statuses; expect coverage warnings (untagged share) in headline"),
        ("Which customers are chasing us for a reply?",
         "[S] chasing = asks_by_them / acknowledged-but-not-answered — AR-relevant signals"),
        ("How many payment follow-up requests came in the last 7 days?",
         "Payment Follow-up topic volume — the AR/collections core topic"),
        ("List stock enquiries from customers this week.",
         "[T] keyword+day-band list shape; harvests customer names as they appear in messages"),
        ("Show me requests we have not answered in the last 10 days.",
         "no-response status set (request-state machine: answered/acked-only/no-response)"),
        ("Which groups are the most active in the last 7 days?",
         "group activity rank — which groups carry the traffic (feeds AR group focus)"),
        ("What was discussed about order dispatch yesterday?",
         "Order Dispatch topic; day-band; verify topical filtering"),
        ("Who shared tax invoices this week?",
         "Tax Invoice/E-Invoice Sharing topic — invoice chatter = AR evidence surface"),
        ("Which customers are slow to reply to us?",
         "[S] response-time claim — VERIFY (unverified capability)"),
        ("List the customers who sent messages in the last 7 days.",
         "customer-alias harvest for AR identity/CLARIFY queries (short/firm names as typed)"),
    ],
}

total = 0
for cat, qs in SECTIONS.items():
    for q, _ in qs:
        m = PLACEHOLDER.search(q)
        if m:
            print(f"PLACEHOLDER TOKEN in probe: {q!r} ({m.group(0)})")
            sys.exit(1)
        total += 1

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Chat Queries"
ws.append(["Query", "Expected Response", "Remarks", "Expected Tool", "Expected Behavior"])
bold = Font(bold=True)
r = 2
for cat, qs in SECTIONS.items():
    c = ws.cell(row=r, column=1, value=cat)
    c.font = bold
    r += 1
    for q, remark in qs:
        ws.cell(row=r, column=1, value=q)
        ws.cell(row=r, column=3, value=remark)
        ws.cell(row=r, column=4, value="RECON")   # no user-visible tools (backend engine)
        ws.cell(row=r, column=5, value="RECON")   # not graded — recon only
        r += 1
wb.save(OUT)
print(f"WROTE {OUT} | recon_queries={total}")