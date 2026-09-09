#!/usr/bin/env python3
"""Rebuild accounts/collections/queries.xlsx with an expected_tool / expected_behavior
column, derived from the REAL deployed system prompt
(code=collection_and_account_receivables). The agent is READ-ONLY with 4 tools and
explicitly marks ageing/overdue/paid-unpaid/reconciliation as UNSUPPORTED (no tool call).

Tool vocabulary:
  search_customers_master  - resolve a named customer / mobile / code
  getCustomerAccountData   - invoice (filter INV) / payment (filter PYMNT) txns, account totals
  getCustomerAnalytics     - outstanding balances, rankings, account metrics, top-N outstanding
  get_sales                - realized-sales summaries + customer rankings

expected_behavior values:
  TOOL:<tool>   -> agent SHOULD call this tool (may chain search_customers_master first)
  NO_TOOL       -> unsupported per system prompt: agent should make NO data-tool call and say "cannot"
"""
import openpyxl
from openpyxl.styles import Font
from pathlib import Path

OUT = Path("/home/sumit/AgentWork/eval-dashboard/accounts/collections/queries.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

# (category, [(query, expected_behavior), ...])
SECTIONS = {
    "Basic Retrieval": [
        ("Show me all customers with outstanding payments.", "TOOL:getCustomerAnalytics"),        # outstanding ranking
        ("What is our total outstanding receivables?", "TOOL:getCustomerAnalytics"),               # all-time outstanding
        ("Show me all unpaid invoices.", "NO_TOOL"),                                               # unpaid status = unsupported
        ("Which customers currently have overdue payments?", "NO_TOOL"),                          # overdue = unsupported
        ("Show me the latest payment received from each customer.", "TOOL:getCustomerAccountData"),# PYMNT txns (per-customer unsupported but tool is the right call)
        ("Give me the outstanding amount customer-wise.", "TOOL:getCustomerAnalytics"),
        ("Show me all invoices that are still partially unpaid.", "NO_TOOL"),                      # paid/unpaid status = unsupported
        ("Which customers have no outstanding balance?", "TOOL:getCustomerAnalytics"),
    ],
    "Overdue & Aging": [
        ("Show me invoices overdue by more than 30 days.", "NO_TOOL"),
        ("Which customers have payments overdue by more than 60 days?", "NO_TOOL"),
        ("Give me the receivables aging breakup.", "NO_TOOL"),                                      # ageing = unsupported
        ("How much money is overdue by more than 90 days?", "NO_TOOL"),
        ("Which customers have the oldest outstanding invoices?", "TOOL:getCustomerAnalytics"),    # outstanding ranking (age unsupported, but outstanding tool ok)
        ("Show me the top 10 overdue accounts.", "NO_TOOL"),                                       # overdue = unsupported
        ("Which overdue invoices have the highest amounts?", "NO_TOOL"),
        ("Are there any invoices that have been overdue for more than 6 months?", "NO_TOOL"),
    ],
    "Customer-Specific": [
        ("Show me all outstanding invoices for Radha Agencies Pvt Ltd 787.", "TOOL:getCustomerAnalytics"),     # resolve + analytics
        ("How much does Radha Agencies Pvt Ltd 787 currently owe us?", "TOOL:getCustomerAnalytics"),
        ("When was the last payment received from Radha Agencies Pvt Ltd 787?", "TOOL:getCustomerAccountData"),# PYMNT
        ("Show me the payment history of Radha Agencies Pvt Ltd 787.", "TOOL:getCustomerAccountData"),          # PYMNT
        ("Does Radha Agencies Pvt Ltd 787 have any invoices overdue by more than 60 days?", "NO_TOOL"),
        ("What is the total invoiced amount versus the amount paid by Radha Agencies Pvt Ltd 787?", "TOOL:getCustomerAccountData"), # INV + PYMNT
        ("Has Radha Agencies Pvt Ltd 787 been consistently paying late?", "NO_TOOL"),                          # late/overdue = unsupported
        ("Which invoices from Radha Agencies Pvt Ltd 787 are still pending?", "TOOL:getCustomerAccountData"),   # INV records (pending≈unpaid, but INV txns supported)
    ],
    "High-Value / Priority": [
        ("Which customers owe us the most money?", "TOOL:getCustomerAnalytics"),
        ("Who should our collection team follow up with first?", "NO_TOOL"),                        # prioritization/follow-up = unsupported judgment
        ("Show me high-value customers with overdue payments.", "NO_TOOL"),                        # overdue = unsupported
        ("Which overdue accounts pose the biggest collection risk?", "NO_TOOL"),
        ("Identify customers with both high outstanding amounts and long payment delays.", "NO_TOOL"), # delay = unsupported
        ("Which 5 customers should we prioritize for collection today?", "NO_TOOL"),
        ("Which customers have a large outstanding balance but haven't made any recent payments?", "NO_TOOL"), # payment recency/reconciliation
        ("Highlight accounts where the overdue amount is unusually high.", "NO_TOOL"),
    ],
    "Payment Behavior Analysis": [
        ("Which customers usually pay after their due date?", "NO_TOOL"),
        ("Who are our most consistent late-paying customers?", "NO_TOOL"),
        ("Which customers have improved their payment behavior recently?", "NO_TOOL"),
        ("Which customers have started delaying payments compared with previous months?", "NO_TOOL"),
        ("Show me customers who frequently make partial payments.", "NO_TOOL"),                     # partial-payment status unsupported
        ("Are there customers whose payment behavior is getting worse?", "NO_TOOL"),
        ("Which customers usually clear their dues before the due date?", "NO_TOOL"),
        ("Compare the payment behavior of our top customers.", "NO_TOOL"),
    ],
    "Trend & Comparison": [
        ("How has our total outstanding changed over the last 6 months?", "TOOL:getCustomerAnalytics"), # outstanding (trend caveat: all-time, not date-scoped)
        ("Compare this month's collections with last month's.", "NO_TOOL"),                          # collections = unsupported
        ("Are overdue receivables increasing or decreasing?", "NO_TOOL"),
        ("Show me the monthly outstanding receivables trend.", "TOOL:getCustomerAnalytics"),         # outstanding only
        ("Which month had the highest collections?", "NO_TOOL"),
        ("Compare overdue amounts for the last three months.", "NO_TOOL"),
        ("Has our collection performance improved this quarter?", "NO_TOOL"),
        ("Which customers contributed most to the increase in outstanding receivables?", "TOOL:getCustomerAnalytics"),
    ],
    "Calculation & Reasoning": [
        ("What percentage of our total receivables is currently overdue?", "NO_TOOL"),              # overdue % = unsupported
        ("What percentage of outstanding receivables is more than 90 days old?", "NO_TOOL"),        # ageing = unsupported
        ("What is the average outstanding amount per customer?", "TOOL:getCustomerAnalytics"),
        ("What is the average payment delay across customers?", "NO_TOOL"),                         # delay = unsupported
        ("How much have we collected versus the total amount invoiced?", "NO_TOOL"),                # collected = unsupported
        ("What portion of the total outstanding amount comes from our top 10 customers?", "TOOL:getCustomerAnalytics"), # top-N outstanding
        ("If we exclude invoices overdue by less than 30 days, how much is seriously overdue?", "NO_TOOL"),
        ("Which customers account for most of our overdue receivables?", "NO_TOOL"),                # overdue = unsupported
    ],
    "Exception & Anomaly": [
        ("Find customers who have outstanding balances despite making recent payments.", "NO_TOOL"), # reconciliation
        ("Show invoices where a payment was received but some amount is still outstanding.", "NO_TOOL"),
        ("Are there customers with unusually large increases in outstanding balances?", "TOOL:getCustomerAnalytics"), # outstanding (anomaly caveat)
        ("Find invoices where the payment appears to be significantly delayed.", "NO_TOOL"),
        ("Are there customers with multiple old invoices but very few recent payments?", "NO_TOOL"),
        ("Show accounts where the outstanding amount is high despite a small number of invoices.", "TOOL:getCustomerAnalytics"),
        ("Identify customers whose current outstanding balance is significantly higher than their historical average.", "TOOL:getCustomerAnalytics"),
        ("Are there any unusual receivable patterns that management should know about?", "NO_TOOL"),
    ],
    "Natural / Human-Like (WhatsApp B2B)": [
        ("Who owes us the most right now?", "TOOL:getCustomerAnalytics"),
        ("Who do we need to chase for payment?", "NO_TOOL"),
        ("Any big payments stuck?", "NO_TOOL"),
        ("Which customers are seriously overdue?", "NO_TOOL"),
        ("Show me the customers we should call today.", "NO_TOOL"),
        ("Are our collections getting better?", "NO_TOOL"),
        ("Who keeps paying late?", "NO_TOOL"),
        ("Where is most of our money stuck?", "TOOL:getCustomerAnalytics"),                          # outstanding concentration
        ("Which old invoices are still not cleared?", "NO_TOOL"),                                  # cleared/overdue = unsupported
        ("Tell me where our biggest collection problems are.", "NO_TOOL"),
    ],
    "Multi-Condition": [
        ("Show me customers with more than 10 lakh outstanding and payments overdue by more than 60 days.", "NO_TOOL"), # overdue condition
        ("Which customers have at least three overdue invoices?", "NO_TOOL"),
        ("Show me invoices above 5 lakh that are overdue by more than 30 days.", "NO_TOOL"),
        ("Which customers have outstanding balances above their average monthly purchase value?", "TOOL:getCustomerAnalytics"), # outstanding (purchase-value join caveat)
        ("Show me customers with high outstanding amounts and no payment in the last 30 days.", "NO_TOOL"), # payment recency
        ("Which customers have more than 20 lakh outstanding, with at least one invoice older than 90 days?", "NO_TOOL"), # aging
    ],
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Chat Queries"
ws.append(["Query", "Expected Response", "Remarks", "Expected Tool"])
bold = Font(bold=True)
r = 2
total = 0
supp = 0
for cat, qs in SECTIONS.items():
    c = ws.cell(row=r, column=1, value=cat); c.font = bold
    r += 1
    for q, beh in qs:
        ws.cell(row=r, column=1, value=q)
        ws.cell(row=r, column=4, value=beh)   # Expected Tool col = expected_behavior
        if beh == "NO_TOOL":
            supp += 1
        r += 1
        total += 1
wb.save(OUT)
print(f"WROTE {OUT} | queries={total} | TOOL-expected={total-supp} | NO_TOOL(expected)={supp}")
