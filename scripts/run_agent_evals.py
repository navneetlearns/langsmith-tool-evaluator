#!/usr/bin/env python3
"""Two-turn-aware eval runner for the agent-template accounts (finance, ar-agent).

Reads accounts/<name>/queries.xlsx (Format A + col E expected_behavior), streams each query
through the shared SSE client (fixed parser: full ui-markdown answers + SSE error capture),
records per query: response, tool_calls, status_sequence, ui_payload_type, suggestions, timing,
error, plus the xlsx labels (category, expected_tool, expected_behavior). Incremental JSONL,
versioned runs, --resume N, NO retries on stream failures (HEART).

Clarify handling: the finance_clarification wire shape is UNVERIFIED (2026-09-18; the early
empty responses were 402 quota errors, now captured as errors). This runner records each turn
faithfully and flags a suspected clarify park (empty response + no error + no tools + fast) as
possible_clarify for the readout; it does NOT fabricate resume turns.
"""
import argparse, json, sys, time, uuid
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from copilot_query_pipeline import (load_account_config, CopilotAuth, CopilotClient,
                                    get_next_version, load_manifest, save_manifest,
                                    update_manifest)

def parse_xlsx(account: str):
    """Format A parser + col E (expected_behavior). Bold col A = category header."""
    wb = openpyxl.load_workbook(f"accounts/{account}/queries.xlsx")
    ws = wb["Chat Queries"] if "Chat Queries" in wb.sheetnames else wb[wb.sheetnames[0]]
    queries, cat = [], "General"
    for idx in range(2, ws.max_row + 1):
        q = ws.cell(idx, 1).value
        if not q:
            continue
        is_bold = bool(ws.cell(idx, 1).font and ws.cell(idx, 1).font.bold)
        if is_bold:
            cat = str(q).strip()
            continue
        qd = {"query": str(q).strip(), "category": cat,
              "expected_tool": ws.cell(idx, 4).value,
              "expected_behavior": ws.cell(idx, 5).value}
        queries.append(qd)
    wb.close()
    return queries

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--account", required=True)
    ap.add_argument("--run", type=int, default=None)
    ap.add_argument("--resume", type=int, default=None)
    args = ap.parse_args()
    account = args.account
    queries = parse_xlsx(account)
    print(f"[{account}] parsed {len(queries)} queries from xlsx")

    cfg = load_account_config(account)
    runs_dir = Path("accounts") / account / "runs"
    manifest_file = runs_dir / "manifest.json"
    version = args.run or get_next_version(runs_dir)
    out_file = runs_dir / f"query_results_v{version}.jsonl"

    completed = {}
    if args.resume:
        resume_file = runs_dir / f"query_results_v{args.resume}.jsonl"
        for line in resume_file.read_text().splitlines() if resume_file.exists() else []:
            try:
                r = json.loads(line)
                completed[r["query_index"]] = r
            except Exception:
                pass
        print(f"[{account}] resume: {len(completed)} already-completed rows loaded from v{args.resume}")

    auth = CopilotAuth(cfg["phone"], cfg["base_url"]); auth.login()
    client = CopilotClient(auth, cfg)

    results, t0_all = [], time.time()
    for i, q in enumerate(queries, 1):
        qidx = i
        if qidx in completed:
            results.append(completed[qidx])
            print(f"[{account}] q{qidx}/{len(queries)} skipped (resume)")
            continue
        tid = str(uuid.uuid4())
        t0 = time.time()
        init_ok = client.init_thread(tid)
        if not init_ok:
            rec = {"query_index": qidx, "query": q["query"], "category": q["category"],
                   "expected_tool": q["expected_tool"], "expected_behavior": q["expected_behavior"],
                   "response": "", "tool_calls": [], "status_sequence": [], "ui_payload_type": None,
                   "suggestions": [], "response_time_seconds": round(time.time() - t0, 2),
                   "error": "Thread init failed", "thread_id": tid}
        else:
            res = client.stream_query(tid, q["query"])
            resp = (res.get("response") or "").strip()
            rec = {"query_index": qidx, "query": q["query"], "category": q["category"],
                   "expected_tool": q["expected_tool"], "expected_behavior": q["expected_behavior"],
                   "response": resp, "tool_calls": res.get("tool_calls", []),
                   "status_sequence": res.get("status_sequence", [])[:60],
                   "ui_payload_type": res.get("ui_payload_type"),
                   "suggestions": res.get("suggestions", []),
                   "response_time_seconds": res.get("response_time_seconds"),
                   "error": res.get("error"), "thread_id": tid}
            # suspected clarify park: empty + no error + no tools + fast
            if (not resp and not rec["error"] and not rec["tool_calls"]
                    and (rec.get("response_time_seconds") or 0) < 5):
                rec["possible_clarify"] = True
        results.append(rec)
        with open(out_file, "a") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        tools = [t.get("tool") for t in rec["tool_calls"]]
        print(f"[{account}] q{qidx}/{len(queries)} {rec.get('response_time_seconds',0):6.1f}s "
              f"err={str(rec['error'])[:60]} tools={tools} clarify={rec.get('possible_clarify', False)}")

    success = sum(1 for r in results if not r.get("error") and (r.get("response") or "").strip())
    failed = len(results) - success
    avg = sum(r.get("response_time_seconds", 0) for r in results) / max(len(results), 1)
    update_manifest(manifest_file, version, out_file, len(results), success, failed, avg)
    print(f"[{account}] v{version} DONE: {success}/{len(results)} ok, {failed} failed, avg {avg:.1f}s -> {out_file}")

if __name__ == "__main__":
    main()