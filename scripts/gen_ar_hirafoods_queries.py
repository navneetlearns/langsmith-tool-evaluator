#!/usr/bin/env python3
"""Rebuild accounts/hirafoods-ar/queries.xlsx — the 32-row HiraFoods AR eval set.

Source of the set: accounts/hirafoods-ar/QUERY_SET_DRAFT_v1.md (approved by the user
2026-10-09; "the draft 32 queries you prepared yesterday is fine - those were customer and
invoice anchored"). The DRAFT is the spec; THIS generator owns the xlsx (never hand-edit it).

expected_tool labels come from the REAL wire surface observed on this workspace, not from a
nominal tool family (AR v1 scored 0/54 strict precisely because its labels came from a family
that never fired):
  get_receivables | ar_payments_reported | ar_promises | search_customers_master |
  getCustomerAnalytics | get_collections | getCustomerAccountData | search_threads |
  list_orders | ASK_BACK | NO_TOOL
Labels marked "[label unverified in gw]" fired no observed tool on the 2026-10-09 gate probe;
they are the closest surface match and drift is reported by the analyzer, never hidden.

expected_behavior: ANSWER | CLARIFY | REFUSE | NEAREST_SHAPE (graded marginal, not fail).
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "accounts/hirafoods-ar/queries.xlsx"
ENTITIES = json.loads((ROOT / "accounts/hirafoods-ar/entities.json").read_text())

# (category, [(query, expected_tool, expected_behavior, remark), ...]) — draft order, rows 1-32.
SECTIONS = {
    "R1 - Customer-anchored": [
        ("Lakshmi Agencies ka kitna baaki hai?", "get_receivables", "ANSWER",
         "LEDGER GATE - gate probe 1: total outstanding ₹57.77Cr answered from get_receivables"),
        ("Laxmi Agency ka outstanding kya hai?", "search_customers_master", "CLARIFY",
         "alias vs master; gate probe 4 OBSERVED an honest non-resolve (asked for mobile/code) - a correct clarify is a pass, guessing is the failure"),
        ("Lakshmi ne bola tha INV-4007 ka paid kar diya - confirm hua?", "ar_payments_reported", "ANSWER",
         "claim vs ledger; gate probe 3 CONFIRMED the INV-4007 claim is live and unresolved; hedge + quote"),
        ("Lakshmi ne payment ka kya promise kiya tha?", "ar_promises", "ANSWER",
         "quote 'INV-4011 by 10-Oct' / 'revert by 8 PM'"),
        ("Lakshmi Agencies ke saath kya stuck hai?", "get_receivables", "ANSWER",
         "open-thread summary; INV-4008 godown-confirm"),
        ("kaunsi party ne outstanding acknowledge kiya?", "ar_promises", "ANSWER",
         "list + cited evidence"),
        ("kis-kis ne bola paid, par system me abhi pending hai?", "ar_payments_reported", "ANSWER",
         "gate probe 3: 1 customer (Lakshmi / INV-4007), claim marked unresolved - must NOT be called paid"),
        ("Lakshmi ke group me pichle hafte kya hua?", "search_threads", "ANSWER",
         "[label unverified in gw] group-activity summary + evidence"),
        ("sabse zyada kis party ka payment baaki hai?", "get_receivables", "ANSWER",
         "LEDGER GATE - top-by-outstanding"),
        ("kis party se last 3 din me koi baat nahi hui?", "search_threads", "ANSWER",
         "[label unverified in gw] activity-recency gap"),
    ],
    "R2 - Invoice-anchored": [
        ("INV-4009 ka payment aaya?", "get_receivables", "ANSWER",
         "link sent ₹7,60,000; 'sent, awaiting confirmation' - hedge, never 'paid'"),
        ("INV-4007 settle hua?", "get_receivables", "ANSWER",
         "customer claimed settled; report ledger truth, do not confirm settlement"),
        ("invoice 4008 ka kya hua?", "get_receivables", "ANSWER",
         "godown confirm pending"),
        ("INV-4013 ka status?", "get_receivables", "ANSWER",
         "qty-change request noted (30 cases Diamond Juice Strong)"),
        ("konsi invoices abhi pending hain?", "get_receivables", "ANSWER",
         "gate probe 2: 14,592 pending invoices answered from get_receivables"),
        ("is hafte dispatch hue invoices?", "getCustomerAccountData", "ANSWER",
         "[label unverified in gw] dispatch-window list"),
        ("kis invoice ke liye party ne bola paid?", "ar_payments_reported", "ANSWER",
         "claim rows; hedge unconfirmed"),
        ("kis invoice ka payment kab tak aayega?", "ar_promises", "ANSWER",
         "PLAN DEVIATION: the draft expected REFUSE (the Zainab B2 boundary). On HiraFoods the invoice-x-promise JOIN ANSWERED (gate probe 6) with a names-don't-match hedge - graded ANSWER-with-hedge, not REFUSE, until a repeat probe contradicts it"),
        ("sabse purani unpaid invoice konsi hai?", "get_receivables", "ANSWER",
         "LEDGER GATE - ageing/oldest"),
        ("INV-4009 tak kitna payment aana hai?", "get_receivables", "ANSWER",
         "LEDGER GATE - amount due"),
    ],
    "D - Derivations": [
        ("total outstanding kitna hai?", "get_receivables", "ANSWER",
         "gate probe 1 LIVE: ₹57,77,21,622.19"),
        ("is mahine kitna payment aaya?", "get_collections", "ANSWER",
         "server-calculated period total only, never a visible-page sum"),
        ("aaj kis-kis ne payment ka promise kiya?", "ar_promises", "ANSWER",
         "dated-promise list"),
        ("top 5 party by outstanding", "getCustomerAnalytics", "ANSWER",
         "LEDGER GATE - ranking (never sum a paginated list)"),
        ("outstanding pichle hafte se badha ya ghata?", "get_receivables", "REFUSE",
         "deployed contract lists 'historical as-of receivables' as UNSUPPORTED - honest refusal is the correct answer; no fabricated trend"),
        ("kis invoice par credit terms hain?", "get_receivables", "ANSWER",
         "₹1,460.88 due 20-Oct anchor"),
    ],
    "Controls": [
        ("Laxmi wale ka kya status hai?", "search_customers_master", "CLARIFY",
         "alias vs master; gate probe 5 OBSERVED the agent answering ORDER status (search_customers_master + list_orders) instead of clarifying - expected CLARIFY, drift recorded"),
        ("aur INV-4011?", "ASK_BACK", "CLARIFY",
         "referent-less follow-up - correct behaviour is to ask which customer/invoice"),
        ("Ultra Biscuits Regular ka stock kitna bacha?", "NO_TOOL", "REFUSE",
         "inventory is not AR; honest out-of-scope, no fabrication"),
        ("Lakshmi Agencies ko supply band kar du?", "NO_TOOL", "REFUSE",
         "credit decision is not AR (real entity - no placeholders in this set)"),
        ("Diamond Juice Strong ka rate kya hai?", "NO_TOOL", "REFUSE",
         "product pricing is not AR"),
        ("Lakshmi ke payment delay kyun hote hain?", "NO_TOOL", "NEAREST_SHAPE",
         "reasons-shape capability gap - graded marginal, not fail"),
    ],
}

# ---- anchors asserted against entities.json (real-entity rule) ----
ANCHOR_CUSTOMERS = ["Lakshmi Agencies", "Laxmi Agency", "Laxmi wale"]
ANCHOR_INVOICES = ["INV-4007", "INV-4008", "INV-4009", "INV-4011", "INV-4013"]
PLACEHOLDER = ("ABC", "XYZ", "TEST", "PLACEHOLDER", "EXAMPLE", "DUMMY", " X ")

ent_names = {c["name"].lower() for c in ENTITIES["customers"]}
for c in ENTITIES["customers"]:
    ent_names.update(a.lower() for a in (c.get("aliases") or []))
ent_inv = {i["number"] for i in ENTITIES["invoices"]}

total = 0
for cat, qs in SECTIONS.items():
    for q, tool, beh, remark in qs:
        for tok in PLACEHOLDER:
            if tok in q:
                print(f"PLACEHOLDER TOKEN in query: {q!r}")
                sys.exit(1)
        for a in ANCHOR_CUSTOMERS:
            if a.lower() in q.lower() and not any(a.lower() in n for n in ent_names):
                print(f"ANCHOR CUSTOMER MISSING FROM entities.json: {a!r} (in {q!r})")
                sys.exit(1)
        for inv in ANCHOR_INVOICES:
            if inv in q and inv not in ent_inv:
                print(f"INVOICE ANCHOR MISSING: {inv} (in {q!r})")
                sys.exit(1)
        total += 1

all_rows = [r for qs in SECTIONS.values() for r in qs]
n_clarify = sum(1 for r in all_rows if r[2] == "CLARIFY")
n_refuse = sum(1 for r in all_rows if r[2] == "REFUSE")
n_gap = sum(1 for r in all_rows if r[2] == "NEAREST_SHAPE")
assert total == 32, f"expected 32 rows, built {total}"
assert n_clarify >= 1 and n_refuse >= 1, "need CLARIFY and REFUSE controls"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Chat Queries"
ws.append(["Query", "Expected Response", "Remarks", "Expected Tool", "Expected Behavior"])
bold = Font(bold=True)
r = 2
for cat, qs in SECTIONS.items():
    c = ws.cell(row=r, column=1, value=cat)
    c.font = bold
    r += 1
    for q, tool, beh, remark in qs:
        ws.cell(row=r, column=1, value=q)
        ws.cell(row=r, column=3, value=remark)
        ws.cell(row=r, column=4, value=tool)
        ws.cell(row=r, column=5, value=beh)
        r += 1
wb.save(OUT)
print(f"WROTE {OUT} | queries={total} | ANSWER={total - n_clarify - n_refuse - n_gap} | "
      f"CLARIFY={n_clarify} | REFUSE={n_refuse} | NEAREST_SHAPE={n_gap}")
